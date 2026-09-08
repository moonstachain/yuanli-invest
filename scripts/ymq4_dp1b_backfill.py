#!/usr/bin/env python3
"""YMQ4-DP1-B historical Gold PIT/as-of backfill.

Physical chain:
World Bank/FRED/ALFRED -> immutable raw SHA -> Supabase private S3 ->
source snapshots / source observations -> monthly PIT/as-of panel -> coverage and
future-leakage gate -> runtime.reality_gate_runs.

DP1-B is a data-plane battle only. It MUST NOT execute or tune B2-B7.
"""
from __future__ import annotations

import calendar
import hashlib
import io
import json
import os
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timezone
from typing import Any, Iterable
from zoneinfo import ZoneInfo

CORE_FACTORS = ("gold_usd_oz", "usd", "inflation_yoy", "real_rate")
BATTLE = "YMQ4-DP1-B"
PANEL_ID = "gold_core_monthly_v0.1"
BUCKET = os.getenv("YMQ4_RAW_BUCKET", "ymq4-raw-evidence")
SUPABASE_PROJECT_REF = os.getenv("YMQ4_SUPABASE_PROJECT_REF", "tbmoimbdhsrltvospwpu")
SUPABASE_REGION = os.getenv("YMQ4_SUPABASE_REGION", "us-east-2")
FRED_ENDPOINT = "https://api.stlouisfed.org/fred/series/observations"
WORLD_BANK_GOLD_URLS = (
    "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx",
    "https://thedocs.worldbank.org/en/doc/561011486076393416-0050022017/original/CMOHistoricalDataMonthly.xlsx",
)
REPLAY_WINDOWS = {
    "gold_1978_1979": (date(1978, 1, 1), date(1979, 12, 31)),
    "gfc_2008_2009": (date(2008, 1, 1), date(2009, 12, 31)),
    "covid_2020": (date(2020, 1, 1), date(2020, 12, 31)),
    "rates_2022": (date(2022, 1, 1), date(2022, 12, 31)),
    "ai_2023_2026": (date(2023, 1, 1), date(2026, 12, 31)),
}


def month_end(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


def month_ends(start: date, end: date) -> list[date]:
    """Return calendar month-ends overlapping [start, end]."""
    if end < start:
        return []
    out: list[date] = []
    year, month = start.year, start.month
    while (year, month) <= (end.year, end.month):
        d = month_end(year, month)
        if d >= start and d <= end:
            out.append(d)
        month += 1
        if month == 13:
            year += 1
            month = 1
    return out


def last_complete_month(today: date | None = None) -> date:
    today = today or datetime.now(timezone.utc).date()
    if today.month == 1:
        return date(today.year - 1, 12, 31)
    return month_end(today.year, today.month - 1)


def _as_date(value: str | date) -> date:
    return value if isinstance(value, date) else date.fromisoformat(value)


def latest_on_or_before(rows: Iterable[dict[str, Any]], decision_date: date) -> dict[str, Any] | None:
    """Select latest non-missing observation whose observation date is <= decision."""
    eligible = [
        row
        for row in rows
        if row.get("value") is not None
        and row.get("observation_date")
        and _as_date(row["observation_date"]) <= decision_date
    ]
    if not eligible:
        return None
    return max(eligible, key=lambda row: _as_date(row["observation_date"]))


def _add_months(d: date, months: int) -> date:
    absolute = d.year * 12 + (d.month - 1) + months
    year, month0 = divmod(absolute, 12)
    month = month0 + 1
    return date(year, month, min(d.day, calendar.monthrange(year, month)[1]))


def cpi_yoy_from_same_vintage(rows: Iterable[dict[str, Any]], latest_observation_date: date) -> float:
    """Compute YoY % using two CPI levels from one as-of-vintage response."""
    levels: dict[date, float] = {}
    for row in rows:
        if row.get("value") is None or not row.get("observation_date"):
            continue
        levels[_as_date(row["observation_date"])] = float(row["value"])
    current_key = date(latest_observation_date.year, latest_observation_date.month, 1)
    prior_key = _add_months(current_key, -12)
    current = levels.get(current_key)
    prior = levels.get(prior_key)
    if current is None or prior in (None, 0):
        raise ValueError("same-vintage CPI YoY requires current and exact 12-month-prior levels")
    return (current / prior - 1.0) * 100.0


def measurement_regime(decision_date: date) -> dict[str, str]:
    """Frozen DP1-B measurement regime for a decision date."""
    if decision_date < date(2003, 1, 1):
        return {
            "id": "1978_2002_proxy",
            "usd_series": "DTWEXM",
            "real_rate_method": "DTB3_MINUS_CPI_YOY_ASOF",
        }
    if decision_date < date(2006, 1, 1):
        return {
            "id": "2003_2005_tips_legacy_usd",
            "usd_series": "DTWEXM",
            "real_rate_method": "DFII10",
        }
    return {
        "id": "2006_plus_modern",
        "usd_series": "DTWEXBGS",
        "real_rate_method": "DFII10",
    }


def future_leakage_count(rows: Iterable[dict[str, Any]]) -> int:
    """Count materialized values that were not yet known on their decision date."""
    leaks = 0
    for row in rows:
        decision = row.get("decision_date")
        known = row.get("known_as_of")
        if decision and known and _as_date(known) > _as_date(decision):
            leaks += 1
    return leaks


def coverage_report(
    rows: Iterable[dict[str, Any]],
    replay_windows: dict[str, tuple[date, date]],
    *,
    threshold: float = 0.80,
) -> dict[str, dict[str, Any]]:
    """Compute complete-month coverage, requiring all four core factors."""
    material = list(rows)
    result: dict[str, dict[str, Any]] = {}
    for name, (start, end) in replay_windows.items():
        expected_dates = month_ends(start, end)
        expected = len(expected_dates)
        by_date: dict[date, set[str]] = {d: set() for d in expected_dates}
        for row in material:
            if row.get("value_numeric") is None:
                continue
            d_raw = row.get("decision_date")
            factor = row.get("factor_id")
            if not d_raw or factor not in CORE_FACTORS:
                continue
            d = _as_date(d_raw)
            if d in by_date:
                by_date[d].add(factor)
        complete = sum(1 for factors in by_date.values() if set(CORE_FACTORS).issubset(factors))
        coverage = complete / expected if expected else 0.0
        result[name] = {
            "expected_months": expected,
            "complete_months": complete,
            "coverage": coverage,
            "threshold": threshold,
            "pass": coverage >= threshold,
        }
    return result


def require_env(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        raise RuntimeError(f"required secret/env missing: {name}")
    return value


def redact_api_key(url: str) -> str:
    parsed = urllib.parse.urlsplit(url)
    q = urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
    q2 = [(k, "REDACTED" if k == "api_key" else v) for k, v in q]
    return urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, parsed.path, urllib.parse.urlencode(q2), parsed.fragment))


def request_bytes(
    url: str,
    *,
    headers: dict[str, str] | None = None,
    data: bytes | None = None,
    method: str | None = None,
    attempts: int = 5,
) -> tuple[int, dict[str, str], bytes]:
    headers = {"User-Agent": "YMQ4-DP1-B/0.1 research-data-backfill", **(headers or {})}
    for attempt in range(attempts):
        req = urllib.request.Request(url, headers=headers, data=data, method=method)
        try:
            with urllib.request.urlopen(req, timeout=120) as resp:
                return resp.status, dict(resp.headers.items()), resp.read()
        except urllib.error.HTTPError as exc:
            if exc.code in (429, 500, 502, 503, 504) and attempt + 1 < attempts:
                time.sleep(min(2 ** attempt, 10))
                continue
            raise RuntimeError(f"HTTP {exc.code} for {redact_api_key(url)}") from None
        except urllib.error.URLError:
            if attempt + 1 < attempts:
                time.sleep(min(2 ** attempt, 10))
                continue
            raise RuntimeError(f"network request failed for {redact_api_key(url)}") from None
    raise RuntimeError(f"request exhausted retries for {redact_api_key(url)}")


def fred_today() -> date:
    return datetime.now(timezone.utc).astimezone(ZoneInfo("America/Chicago")).date()


def fred_url(api_key: str, series_id: str, **params: Any) -> str:
    query = {
        "series_id": series_id,
        "api_key": api_key,
        "file_type": "json",
        "limit": 100000,
        **params,
    }
    return FRED_ENDPOINT + "?" + urllib.parse.urlencode(query)


def _parse_fred_rows(raw: bytes) -> list[dict[str, Any]]:
    payload = json.loads(raw)
    rows: list[dict[str, Any]] = []
    for item in payload.get("observations", []):
        if item.get("value") in (None, "."):
            continue
        rows.append({
            "observation_date": item["date"],
            "value": float(item["value"]),
            "realtime_start": item.get("realtime_start"),
            "realtime_end": item.get("realtime_end"),
        })
    return rows


def fetch_fred_market_series(api_key: str, series_id: str, start: date, end: date) -> tuple[str, int, dict[str, str], bytes, list[dict[str, Any]]]:
    url = fred_url(
        api_key,
        series_id,
        observation_start=start.isoformat(),
        observation_end=end.isoformat(),
        output_type=1,
    )
    status, headers, raw = request_bytes(url)
    rows = _parse_fred_rows(raw)
    if not rows:
        raise RuntimeError(f"no non-missing observations returned for {series_id}")
    return url, status, headers, raw, rows


def fetch_cpi_initial_releases(api_key: str, start: date, end: date) -> tuple[str, int, dict[str, str], bytes, list[dict[str, Any]]]:
    url = fred_url(
        api_key,
        "CPIAUCSL",
        observation_start=start.isoformat(),
        observation_end=end.isoformat(),
        output_type=4,
        realtime_start=start.isoformat(),
        realtime_end=fred_today().isoformat(),
    )
    status, headers, raw = request_bytes(url)
    rows = _parse_fred_rows(raw)
    if not rows:
        raise RuntimeError("CPI initial-release query returned no observations")
    return url, status, headers, raw, rows


def fetch_cpi_asof(api_key: str, decision: date) -> tuple[dict[str, Any], bytes]:
    observation_start = _add_months(date(decision.year, decision.month, 1), -15)
    url = fred_url(
        api_key,
        "CPIAUCSL",
        observation_start=observation_start.isoformat(),
        observation_end=decision.isoformat(),
        output_type=1,
        realtime_start=decision.isoformat(),
        realtime_end=decision.isoformat(),
    )
    _, _, raw = request_bytes(url)
    rows = _parse_fred_rows(raw)
    if not rows:
        raise RuntimeError(f"CPI as-of query returned no data for {decision}")
    latest = latest_on_or_before(rows, decision)
    if latest is None:
        raise RuntimeError(f"CPI as-of has no eligible level for {decision}")
    latest_obs = _as_date(latest["observation_date"])
    yoy = cpi_yoy_from_same_vintage(rows, latest_obs)
    return {
        "decision_date": decision.isoformat(),
        "latest_observation_date": latest_obs.isoformat(),
        "value": yoy,
        "request_template": redact_api_key(url),
        "rows": rows,
    }, raw


def parse_world_bank_monthly_gold(raw: bytes) -> list[dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to parse World Bank Pink Sheet") from exc
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    candidates = [ws for ws in wb.worksheets if "month" in ws.title.lower()]
    if not candidates:
        candidates = list(wb.worksheets)
    for ws in candidates:
        gold_col = None
        header_row = None
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True), start=1):
            for col_idx, cell in enumerate(row, start=1):
                if isinstance(cell, str) and cell.strip().lower() == "gold":
                    gold_col = col_idx
                    header_row = row_idx
                    break
            if gold_col:
                break
        if not gold_col or not header_row:
            continue
        out: list[dict[str, Any]] = []
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if gold_col > len(row):
                continue
            value = row[gold_col - 1]
            if value in (None, ""):
                continue
            parsed_date = None
            for cell in row[: min(4, len(row))]:
                parsed_date = parse_month_cell(cell)
                if parsed_date:
                    break
            if parsed_date is None:
                continue
            try:
                numeric = float(value)
            except (TypeError, ValueError):
                continue
            out.append({"observation_date": month_end(parsed_date.year, parsed_date.month).isoformat(), "value": numeric})
        if len(out) >= 400:
            return out
    raise RuntimeError("World Bank workbook did not yield a valid monthly Gold column")


def parse_month_cell(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date().replace(day=1)
    if isinstance(value, date):
        return value.replace(day=1)
    if not isinstance(value, str):
        return None
    text = value.strip()
    if not text:
        return None
    for sep in ("M", "-", "/"):
        if sep in text:
            left, right, *_ = text.replace("/", "-").replace("M", "-").split("-") + [""]
            if left.isdigit() and right.isdigit() and len(left) == 4:
                try:
                    return date(int(left), int(right), 1)
                except ValueError:
                    return None
    for fmt in ("%b-%y", "%b %Y", "%B %Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date().replace(day=1)
        except ValueError:
            pass
    return None


def fetch_world_bank_gold() -> tuple[str, int, dict[str, str], bytes, list[dict[str, Any]]]:
    errors: list[str] = []
    for url in WORLD_BANK_GOLD_URLS:
        try:
            status, headers, raw = request_bytes(url)
            rows = parse_world_bank_monthly_gold(raw)
            return url, status, headers, raw, rows
        except Exception as exc:
            errors.append(f"{url}: {type(exc).__name__}")
    raise RuntimeError("all governed World Bank Pink Sheet URLs failed: " + "; ".join(errors))


def rpc_headers(secret_key: str) -> dict[str, str]:
    if not secret_key.startswith("sb_secret_"):
        raise RuntimeError("YMQ4_SUPABASE_SECRET_KEY must be a modern sb_secret_ key")
    return {"apikey": secret_key, "Content-Type": "application/json"}


def rpc(supabase_url: str, secret_key: str, function_name: str, payload: dict[str, Any]) -> Any:
    url = supabase_url.rstrip("/") + "/rest/v1/rpc/" + function_name
    _, _, raw = request_bytes(url, headers=rpc_headers(secret_key), data=json.dumps(payload).encode(), method="POST")
    return json.loads(raw or b"null")


def s3_client(access_key_id: str, secret_access_key: str):
    try:
        import boto3
        from botocore.config import Config
    except ImportError as exc:
        raise RuntimeError("boto3 is required for the full DP1-B S3 backfill") from exc
    endpoint = f"https://{SUPABASE_PROJECT_REF}.storage.supabase.co/storage/v1/s3"
    return boto3.client(
        "s3",
        region_name=SUPABASE_REGION,
        endpoint_url=endpoint,
        aws_access_key_id=access_key_id,
        aws_secret_access_key=secret_access_key,
        config=Config(s3={"addressing_style": "path"}, signature_version="s3v4"),
    )


def archive_raw_s3(client: Any, source_id: str, raw: bytes, content_type: str, git_sha: str) -> dict[str, str]:
    sha = hashlib.sha256(raw).hexdigest()
    stamp = datetime.now(timezone.utc).strftime("%Y/%m/%d/%H%M%S")
    ext = "xlsx" if "spreadsheet" in content_type or raw[:2] == b"PK" else "json"
    path = f"dp1b/{source_id}/{stamp}-{sha}.{ext}"
    client.put_object(
        Bucket=BUCKET,
        Key=path,
        Body=raw,
        ContentType=content_type,
        Metadata={"sha256": sha, "battle": BATTLE, "runner-commit": git_sha},
    )
    obj = client.get_object(Bucket=BUCKET, Key=path)
    reread = obj["Body"].read()
    if hashlib.sha256(reread).hexdigest() != sha:
        raise RuntimeError(f"S3 SHA readback mismatch for {source_id}")
    if (obj.get("Metadata") or {}).get("sha256") != sha:
        raise RuntimeError(f"S3 metadata SHA mismatch for {source_id}")
    return {"sha256": sha, "path": path}


def normalized_market_observations(series_id: str, rows: list[dict[str, Any]], regime: str) -> list[dict[str, Any]]:
    return [
        {
            "series_id": series_id,
            "value_numeric": row["value"],
            "observation_date": row["observation_date"],
            "release_date": row["observation_date"],
            "vintage_date": row["observation_date"],
            "known_as_of": row["observation_date"],
            "pit_status": "PIT_MARKET_RECONSTRUCTED",
            "measurement_regime": regime,
        }
        for row in rows
    ]


def normalized_cpi_initial(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        release = row.get("realtime_start")
        if not release:
            continue
        if _as_date(release) < _as_date(row["observation_date"]):
            continue
        out.append({
            "series_id": "CPIAUCSL",
            "value_numeric": row["value"],
            "observation_date": row["observation_date"],
            "release_date": release,
            "vintage_date": release,
            "known_as_of": release,
            "pit_status": "PIT_STRICT_INITIAL_RELEASE",
            "measurement_regime": "CPI_INITIAL_RELEASE",
        })
    if not out:
        raise RuntimeError("CPI initial releases could not establish four clocks")
    return out


def ingest_source_batches(
    sb_url: str,
    sb_key: str,
    source_id: str,
    archive: dict[str, str],
    retrieved_at: datetime,
    http_status: int,
    content_type: str,
    request_template: str,
    git_sha: str,
    observations: list[dict[str, Any]],
    batch_size: int = 800,
) -> str:
    batches = [observations[i:i + batch_size] for i in range(0, len(observations), batch_size)] or [[]]
    snapshot_id = None
    for batch in batches:
        result = rpc(sb_url, sb_key, "ymq4_dp1b_ingest_source", {
            "p_source_id": source_id,
            "p_retrieved_at": retrieved_at.isoformat(),
            "p_http_status": http_status,
            "p_content_type": content_type,
            "p_sha256": archive["sha256"],
            "p_storage_bucket": BUCKET,
            "p_storage_path": archive["path"],
            "p_request_template": request_template,
            "p_runner_commit": git_sha,
            "p_observations": batch,
        })
        if not isinstance(result, list) or len(result) != 1:
            raise RuntimeError(f"DP1-B ingest RPC returned unexpected shape for {source_id}")
        snapshot_id = result[0]["snapshot_id"]
    if not snapshot_id:
        raise RuntimeError(f"DP1-B ingest did not return snapshot for {source_id}")
    return snapshot_id


def build_panel(
    decisions: list[date],
    gold_rows: list[dict[str, Any]],
    market: dict[str, list[dict[str, Any]]],
    cpi_asof: dict[str, dict[str, Any]],
    snapshots: dict[str, str],
) -> list[dict[str, Any]]:
    panel: list[dict[str, Any]] = []
    gold_by_month = {row["observation_date"][:7]: row for row in gold_rows}
    for decision in decisions:
        regime = measurement_regime(decision)
        dstr = decision.isoformat()
        gold = gold_by_month.get(dstr[:7])
        usd = latest_on_or_before(market[regime["usd_series"]], decision)
        cpi = cpi_asof.get(dstr)
        rate = latest_on_or_before(market["DTB3" if regime["real_rate_method"].startswith("DTB3") else "DFII10"], decision)
        if gold:
            panel.append({
                "decision_date": dstr,
                "factor_id": "gold_usd_oz",
                "value_numeric": gold["value"],
                "source_series_id": "worldbank_pinksheet_gold",
                "pit_status": "PIT_MARKET_RECONSTRUCTED",
                "known_as_of": dstr,
                "measurement_regime": regime["id"],
                "provenance": {"snapshot_id": snapshots["worldbank_pinksheet_gold"], "source_observation_date": gold["observation_date"]},
            })
        if usd:
            panel.append({
                "decision_date": dstr,
                "factor_id": "usd",
                "value_numeric": usd["value"],
                "source_series_id": regime["usd_series"],
                "pit_status": "PIT_MARKET_RECONSTRUCTED",
                "known_as_of": usd["observation_date"],
                "measurement_regime": regime["id"],
                "provenance": {"snapshot_id": snapshots[regime["usd_series"]], "source_observation_date": usd["observation_date"]},
            })
        if cpi:
            panel.append({
                "decision_date": dstr,
                "factor_id": "inflation_yoy",
                "value_numeric": cpi["value"],
                "source_series_id": "CPIAUCSL",
                "pit_status": "PIT_STRICT_ASOF_DERIVED",
                "known_as_of": dstr,
                "measurement_regime": regime["id"],
                "provenance": {"snapshot_id": snapshots["CPI_ASOF_BUNDLE"], "latest_cpi_observation": cpi["latest_observation_date"], "asof_date": dstr, "transform": "same_vintage_yoy"},
            })
        if rate:
            if regime["real_rate_method"].startswith("DTB3"):
                if cpi:
                    value = rate["value"] - cpi["value"]
                    panel.append({
                        "decision_date": dstr,
                        "factor_id": "real_rate",
                        "value_numeric": value,
                        "source_series_id": "DTB3-CPIAUCSL_YOY",
                        "pit_status": "PIT_DERIVED_PROXY",
                        "known_as_of": dstr,
                        "measurement_regime": regime["id"],
                        "provenance": {"rate_snapshot_id": snapshots["DTB3"], "cpi_snapshot_id": snapshots["CPI_ASOF_BUNDLE"], "rate_observation_date": rate["observation_date"], "formula": "DTB3 - CPI_YOY_ASOF"},
                    })
            else:
                panel.append({
                    "decision_date": dstr,
                    "factor_id": "real_rate",
                    "value_numeric": rate["value"],
                    "source_series_id": "DFII10",
                    "pit_status": "PIT_MARKET_RECONSTRUCTED",
                    "known_as_of": rate["observation_date"],
                    "measurement_regime": regime["id"],
                    "provenance": {"snapshot_id": snapshots["DFII10"], "source_observation_date": rate["observation_date"]},
                })
    return panel


def clamp_replay_windows(end: date) -> dict[str, tuple[date, date]]:
    out = {}
    for name, (start, stop) in REPLAY_WINDOWS.items():
        if start > end:
            continue
        out[name] = (start, min(stop, end))
    return out


def main() -> int:
    started = datetime.now(timezone.utc)
    api_key = require_env("FRED_API_KEY")
    sb_url = require_env("SUPABASE_URL")
    sb_key = require_env("YMQ4_SUPABASE_SECRET_KEY")
    s3_access = require_env("YMQ4_SUPABASE_S3_ACCESS_KEY_ID")
    s3_secret = require_env("YMQ4_SUPABASE_S3_SECRET_ACCESS_KEY")
    git_sha = os.getenv("GITHUB_SHA", "LOCAL")
    expected_url = f"https://{SUPABASE_PROJECT_REF}.supabase.co"
    if sb_url.rstrip("/") != expected_url:
        raise RuntimeError("SUPABASE_URL project-ref mismatch")

    end = last_complete_month()
    decisions = month_ends(date(1978, 1, 1), end)
    storage = s3_client(s3_access, s3_secret)
    storage.head_bucket(Bucket=BUCKET)
    snapshots: dict[str, str] = {}
    source_receipts: dict[str, dict[str, Any]] = {}
    market: dict[str, list[dict[str, Any]]] = {}

    # World Bank Gold monthly reconstruction.
    wb_url, wb_status, wb_headers, wb_raw, gold_rows = fetch_world_bank_gold()
    retrieved = datetime.now(timezone.utc)
    wb_archive = archive_raw_s3(storage, "worldbank_pinksheet_gold", wb_raw, wb_headers.get("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), git_sha)
    snapshots["worldbank_pinksheet_gold"] = ingest_source_batches(
        sb_url, sb_key, "worldbank_pinksheet_gold", wb_archive, retrieved, wb_status,
        wb_headers.get("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), wb_url, git_sha,
        normalized_market_observations("WORLD_BANK_GOLD_MONTHLY", gold_rows, "WORLD_BANK_MONTHLY_RECONSTRUCTION"),
    )
    source_receipts["worldbank_pinksheet_gold"] = {**wb_archive, "rows": len(gold_rows), "url": wb_url, "snapshot_id": snapshots["worldbank_pinksheet_gold"]}

    # FRED market series: deliberately reconstructed, never silently strict-vintage.
    market_specs = {
        "DTB3": (date(1976, 1, 1), end, "fred_dtb3"),
        "DTWEXM": (date(1973, 1, 1), min(end, date(2019, 12, 31)), "fred_dtwexm"),
        "DTWEXBGS": (date(2006, 1, 1), end, "fred_dtwexbgs"),
        "DFII10": (date(2003, 1, 1), end, "fred_dfii10"),
    }
    for series_id, (start, stop, source_id) in market_specs.items():
        url, status, headers, raw, rows = fetch_fred_market_series(api_key, series_id, start, stop)
        market[series_id] = rows
        retrieved = datetime.now(timezone.utc)
        archive = archive_raw_s3(storage, source_id, raw, headers.get("Content-Type", "application/json"), git_sha)
        snapshots[series_id] = ingest_source_batches(
            sb_url, sb_key, source_id, archive, retrieved, status, headers.get("Content-Type", "application/json"),
            redact_api_key(url), git_sha, normalized_market_observations(series_id, rows, "MARKET_RECONSTRUCTION"),
        )
        source_receipts[source_id] = {**archive, "rows": len(rows), "snapshot_id": snapshots[series_id]}

    # CPI initial-release ledger.
    cpi_start = date(1976, 1, 1)
    cpi_url, cpi_status, cpi_headers, cpi_raw, cpi_initial_rows = fetch_cpi_initial_releases(api_key, cpi_start, end)
    cpi_archive = archive_raw_s3(storage, "fred_cpiaucsl_initial", cpi_raw, cpi_headers.get("Content-Type", "application/json"), git_sha)
    cpi_initial_norm = normalized_cpi_initial(cpi_initial_rows)
    snapshots["CPI_INITIAL"] = ingest_source_batches(
        sb_url, sb_key, "fred_cpiaucsl", cpi_archive, datetime.now(timezone.utc), cpi_status,
        cpi_headers.get("Content-Type", "application/json"), redact_api_key(cpi_url), git_sha, cpi_initial_norm,
    )
    source_receipts["fred_cpiaucsl_initial"] = {**cpi_archive, "rows": len(cpi_initial_norm), "snapshot_id": snapshots["CPI_INITIAL"]}

    # Strict monthly same-vintage CPI state. Each provider response is captured in
    # one immutable bundle so every derived YoY can be audited without 584 S3 objects.
    cpi_asof: dict[str, dict[str, Any]] = {}
    raw_bundle: dict[str, Any] = {"battle": BATTLE, "series_id": "CPIAUCSL", "queries": {}}
    for idx, decision in enumerate(decisions):
        state, raw = fetch_cpi_asof(api_key, decision)
        cpi_asof[decision.isoformat()] = {k: v for k, v in state.items() if k not in ("rows", "request_template")}
        raw_bundle["queries"][decision.isoformat()] = json.loads(raw)
        # Stay below common FRED per-minute limits while retaining retry headroom.
        if idx + 1 < len(decisions):
            time.sleep(0.55)
    bundle_bytes = json.dumps(raw_bundle, separators=(",", ":"), sort_keys=True).encode()
    cpi_asof_archive = archive_raw_s3(storage, "fred_cpiaucsl_asof_bundle", bundle_bytes, "application/json", git_sha)
    snapshots["CPI_ASOF_BUNDLE"] = ingest_source_batches(
        sb_url, sb_key, "fred_cpiaucsl", cpi_asof_archive, datetime.now(timezone.utc), 200,
        "application/json", "FRED/ALFRED CPIAUCSL monthly same-day as-of bundle; API key redacted", git_sha, [],
    )
    source_receipts["fred_cpiaucsl_asof_bundle"] = {**cpi_asof_archive, "rows": len(cpi_asof), "snapshot_id": snapshots["CPI_ASOF_BUNDLE"]}

    panel = build_panel(decisions, gold_rows, market, cpi_asof, snapshots)
    leaks = future_leakage_count(panel)
    if leaks:
        raise RuntimeError(f"future leakage detected before database write: {leaks}")

    for i in range(0, len(panel), 700):
        result = rpc(sb_url, sb_key, "ymq4_dp1b_upsert_panel", {
            "p_panel_id": PANEL_ID,
            "p_rows": panel[i:i + 700],
            "p_runner_commit": git_sha,
        })
        if not isinstance(result, int):
            raise RuntimeError("panel upsert RPC returned unexpected shape")

    readback = rpc(sb_url, sb_key, "ymq4_dp1b_readback", {"p_panel_id": PANEL_ID})
    if not isinstance(readback, dict):
        raise RuntimeError("panel readback RPC returned unexpected shape")
    if int(readback.get("future_leakage", -1)) != 0:
        raise RuntimeError("database readback found future leakage")
    if int(readback.get("total_rows", -1)) != len(panel):
        raise RuntimeError(f"panel row-count readback mismatch: {readback.get('total_rows')} != {len(panel)}")

    windows = clamp_replay_windows(end)
    coverage = coverage_report(panel, windows, threshold=0.80)
    all_windows_pass = bool(coverage) and all(v["pass"] for v in coverage.values())
    status = "DP1B_CORE_BACKFILL_PASS" if all_windows_pass else "PARTIAL_RESEARCH_ONLY"
    factor_counts = {factor: sum(1 for r in panel if r["factor_id"] == factor) for factor in CORE_FACTORS}
    regime_counts: dict[str, int] = {}
    for row in panel:
        regime_counts[row["measurement_regime"]] = regime_counts.get(row["measurement_regime"], 0) + 1

    receipt = {
        "battle": BATTLE,
        "status": status,
        "git_sha": git_sha,
        "panel_id": PANEL_ID,
        "period": {"start": decisions[0].isoformat(), "end": decisions[-1].isoformat(), "months": len(decisions)},
        "panel_rows": len(panel),
        "factor_counts": factor_counts,
        "future_leakage": 0,
        "coverage": coverage,
        "measurement_regime_counts": regime_counts,
        "sources": source_receipts,
        "database_readback": readback,
        "constraints": {"coverage_threshold": 0.80, "future_leakage_tolerance": 0, "b2_b7_executed": False},
        "started_at": started.isoformat(),
        "completed_at": datetime.now(timezone.utc).isoformat(),
    }
    gate_id = rpc(sb_url, sb_key, "ymq4_dp1b_record_gate", {
        "p_git_sha": git_sha,
        "p_started_at": receipt["started_at"],
        "p_completed_at": receipt["completed_at"],
        "p_gate_status": status,
        "p_receipt": receipt,
    })
    receipt["reality_gate_run_id"] = gate_id
    print(json.dumps(receipt, indent=2, sort_keys=True))
    return 0 if status == "DP1B_CORE_BACKFILL_PASS" else 2


def run() -> int:
    try:
        return main()
    except Exception as exc:
        message = str(exc)
        for name in ("FRED_API_KEY", "YMQ4_SUPABASE_SECRET_KEY", "YMQ4_SUPABASE_S3_ACCESS_KEY_ID", "YMQ4_SUPABASE_S3_SECRET_ACCESS_KEY"):
            value = os.getenv(name, "").strip()
            if value:
                for variant in (value, urllib.parse.quote(value, safe=""), urllib.parse.quote_plus(value)):
                    message = message.replace(variant, "REDACTED")
        print(json.dumps({"battle": BATTLE, "status": "FAIL_CLOSED", "error_type": type(exc).__name__, "error": message}, indent=2))
        return 1


if __name__ == "__main__":
    raise SystemExit(run())
